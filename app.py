
from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file
import sqlite3, os, secrets, csv, io, threading, time, re
import openpyxl
from dotenv import load_dotenv
load_dotenv(os.path.join(os.path.dirname(__file__), '.env'))
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
import arabic_reshaper
from bidi.algorithm import get_display
from email_service import send_email, is_configured, fetch_inbox
from datetime import datetime, timedelta

APP = Flask(__name__)
APP.secret_key = os.environ.get("SECRET_KEY", "change-this-secret-key")
DB = os.path.join(os.path.dirname(__file__), "requests.db")

REFERENCES = ["قوى","التأمينات الاجتماعية","مدد","وزارة الموارد الاجتماعية","وزارة التجارة","مقيم","مبادرة","فضائات","التأمين الطبي","سبل","المركز السعودي للأعمال","الغرفة التجارية","الجوازات"]
CITIES = ["الرياض","جدة","مكة المكرمة","المدينة المنورة","الدمام","الخبر","الطائف","تبوك","أبها","خميس مشيط","حائل","جازان","نجران","الباحة","سكاكا","عرعر","بريدة","عنيزة","ينبع","الأحساء","الجبيل","القطيف","رابغ","الخرج","الرس","المجمعة","وغير ذلك"]

def db():
    c=sqlite3.connect(DB)
    c.row_factory=sqlite3.Row
    return c

def init_db():
    c=db()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS users(
      id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, username TEXT UNIQUE NOT NULL,
      password TEXT NOT NULL, email TEXT, role TEXT NOT NULL DEFAULT 'employee', active INTEGER DEFAULT 1
    );
    CREATE TABLE IF NOT EXISTS user_permissions(
      user_id INTEGER PRIMARY KEY, view_requests INTEGER DEFAULT 1, create_request INTEGER DEFAULT 1,
      manage_companies INTEGER DEFAULT 0, reports INTEGER DEFAULT 0, completed INTEGER DEFAULT 1,
      in_progress INTEGER DEFAULT 1, ten_days INTEGER DEFAULT 0, users_admin INTEGER DEFAULT 0,
      email_templates INTEGER DEFAULT 0, audit INTEGER DEFAULT 0, email_integration INTEGER DEFAULT 0,
      email_settings INTEGER DEFAULT 0,
      FOREIGN KEY(user_id) REFERENCES users(id)
    );
    CREATE TABLE IF NOT EXISTS companies(
      id INTEGER PRIMARY KEY AUTOINCREMENT, cr TEXT UNIQUE NOT NULL, unified TEXT NOT NULL, name TEXT NOT NULL, active INTEGER DEFAULT 1
    );
    CREATE TABLE IF NOT EXISTS requests(
      id INTEGER PRIMARY KEY AUTOINCREMENT, serial TEXT UNIQUE NOT NULL, request_no TEXT NOT NULL,
      request_type TEXT NOT NULL, reference TEXT NOT NULL, city TEXT NOT NULL,
      company_id INTEGER, created_by INTEGER, created_at TEXT NOT NULL,
      status TEXT NOT NULL DEFAULT 'قيد التنفيذ', last_followup TEXT, close_requested_at TEXT,
      closed_at TEXT, closed_by INTEGER, close_note TEXT,
      FOREIGN KEY(company_id) REFERENCES companies(id)
    );
    CREATE TABLE IF NOT EXISTS followups(
      id INTEGER PRIMARY KEY AUTOINCREMENT, request_id INTEGER, user_id INTEGER, note TEXT NOT NULL,
      created_at TEXT NOT NULL, FOREIGN KEY(request_id) REFERENCES requests(id)
    );
    CREATE TABLE IF NOT EXISTS audit_logs(
      id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, action TEXT NOT NULL,
      entity TEXT, entity_id INTEGER, details TEXT, created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS email_templates(
      id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT UNIQUE NOT NULL, subject TEXT NOT NULL, body TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS email_messages(
      id INTEGER PRIMARY KEY AUTOINCREMENT, request_id INTEGER, direction TEXT NOT NULL,
      sender TEXT, recipients TEXT, subject TEXT, body TEXT, message_id TEXT UNIQUE,
      received_at TEXT NOT NULL, FOREIGN KEY(request_id) REFERENCES requests(id)
    );
    """)
    cols=[r["name"] for r in c.execute("PRAGMA table_info(user_permissions)").fetchall()]
    if "email_settings" not in cols: c.execute("ALTER TABLE user_permissions ADD COLUMN email_settings INTEGER DEFAULT 0")
    if not c.execute("SELECT 1 FROM users LIMIT 1").fetchone():
        c.execute("INSERT INTO users(name,username,password,email,role) VALUES(?,?,?,?,?)",
                  ("مدير النظام","admin","admin123","","super_admin"))
    templates=[
      ("followup","تمت متابعة الطلب {{request_no}}","تمت متابعة الطلب رقم {{request_no}} بواسطة {{user}} بتاريخ {{date}}.\n\nتفاصيل المتابعة:\n{{note}}"),
      ("close_request","طلب اعتماد إنهاء الطلب {{request_no}}","الموظف {{user}} طلب اعتماد إنهاء الطلب رقم {{request_no}} بتاريخ {{date}}.\n\nالملاحظات:\n{{note}}"),
      ("closed","تم اعتماد إنهاء الطلب {{request_no}}","تم اعتماد إنهاء الطلب رقم {{request_no}} بواسطة {{manager}} بتاريخ {{date}}."),
      ("rejected","تم رفض إنهاء الطلب {{request_no}}","تم رفض طلب إنهاء الطلب رقم {{request_no}} بواسطة {{manager}} بتاريخ {{date}}.\\n\\nسبب الرفض:\\n{{note}}")
    ]
    for t in templates:
        c.execute("INSERT OR IGNORE INTO email_templates(code,subject,body) VALUES(?,?,?)",t)
    for u in c.execute("SELECT id,role FROM users").fetchall():
        if not c.execute("SELECT 1 FROM user_permissions WHERE user_id=?",(u["id"],)).fetchone():
            c.execute("""INSERT INTO user_permissions(user_id,view_requests,create_request,manage_companies,reports,completed,in_progress,ten_days,users_admin,email_templates,audit,email_integration,email_settings)
                         VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(u["id"],1,1,1 if u["role"]=="super_admin" else 0,1 if u["role"] in ("hr","manager","super_admin") else 0,1,1,1 if u["role"] in ("manager","super_admin") else 0,1 if u["role"]=="super_admin" else 0,1 if u["role"]=="super_admin" else 0,1,1 if u["role"]=="super_admin" else 0,1 if u["role"]=="super_admin" else 0))
    c.commit(); c.close()

def now(): return datetime.now().strftime("%Y-%m-%d %H:%M:%S")
def user(): return db().execute("SELECT * FROM users WHERE id=?",(session.get("uid"),)).fetchone()

@APP.context_processor
def inject_current_user():
    return {"current_user": user(), "has_permission": has_permission}

def audit(action, entity="", entity_id=None, details=""):
    c=db(); c.execute("INSERT INTO audit_logs(user_id,action,entity,entity_id,details,created_at) VALUES(?,?,?,?,?,?)",
      (session.get("uid"),action,entity,entity_id,details,now())); c.commit(); c.close()

def render_email(code, data):
    c=db(); t=c.execute("SELECT * FROM email_templates WHERE code=?",(code,)).fetchone(); c.close()
    if not t: return "", ""
    s=t["subject"]; b=t["body"]
    for k,v in data.items():
        s=s.replace("{{"+k+"}}",str(v)); b=b.replace("{{"+k+"}}",str(v))
    return s,b

def _email_subject(subject,data):
    serial=data.get('serial')
    return f'[REQ:{serial}] {subject}' if serial else subject

def notify_managers(code,data):
    subject,body=render_email(code,data)
    subject=_email_subject(subject,data)
    c=db()
    managers=c.execute("SELECT email FROM users WHERE active=1 AND role IN ('manager','super_admin') AND email<>''").fetchall()
    recipients=[x["email"] for x in managers]
    c.close()
    try:
        ok, detail=send_email(recipients, subject, body)
    except Exception as exc:
        ok, detail=False, f"فشل إرسال الإشعار: {exc}"
    audit("إرسال إشعار بريد", "email", None, f"إلى: {', '.join(recipients)}\\nالنتيجة: {detail}\\nالموضوع: {subject}")
    return ok, detail

def notify_user(email, code, data):
    if not email:
        return False, "لا يوجد بريد للمستخدم"
    subject,body=render_email(code,data)
    subject=_email_subject(subject,data)
    try:
        ok, detail=send_email([email], subject, body)
    except Exception as exc:
        ok, detail=False, f"فشل إرسال البريد: {exc}"
    audit("إرسال بريد للمستخدم", "email", None, f"إلى: {email}\\nالنتيجة: {detail}\\nالموضوع: {subject}")
    return ok, detail

@APP.route("/")
def home():
    if not session.get("uid"): return redirect(url_for("login"))
    c=db()
    total=c.execute("SELECT COUNT(*) n FROM requests").fetchone()["n"]
    active=c.execute("SELECT COUNT(*) n FROM requests WHERE status='قيد التنفيذ'").fetchone()["n"]
    closed=c.execute("SELECT COUNT(*) n FROM requests WHERE status='منتهي'").fetchone()["n"]
    pending=c.execute("SELECT COUNT(*) n FROM requests WHERE status='بانتظار اعتماد الإنهاء'").fetchone()["n"]
    cutoff=(datetime.now()-timedelta(days=10)).strftime("%Y-%m-%d %H:%M:%S")
    overdue=c.execute("SELECT COUNT(*) n FROM requests WHERE status='قيد التنفيذ' AND COALESCE(last_followup,created_at) <= ?",(cutoff,)).fetchone()["n"]
    recent=c.execute("""SELECT r.*,u.name creator,co.cr,co.name company_name
                       FROM requests r LEFT JOIN users u ON u.id=r.created_by
                       LEFT JOIN companies co ON co.id=r.company_id
                       ORDER BY r.id DESC LIMIT 8""").fetchall()
    c.close()
    return render_template("dashboard.html", title="لوحة التحكم", total=total,active=active,closed=closed,pending=pending,overdue=overdue,recent=recent)

@APP.route("/login", methods=["GET","POST"])
def login():
    if request.method=="POST":
        c=db(); u=c.execute("SELECT * FROM users WHERE username=? AND password=? AND active=1",
                            (request.form["username"],request.form["password"])).fetchone(); c.close()
        if u:
            session["uid"]=u["id"]; return redirect(url_for("home"))
        flash("بيانات الدخول غير صحيحة","error")
    return render_template("login.html", title="تسجيل الدخول")

@APP.route("/logout")
def logout(): session.clear(); return redirect(url_for("login"))

def has_permission(name):
    u=user()
    if not u: return False
    if u["role"]=="super_admin": return True
    c=db(); p=c.execute("SELECT * FROM user_permissions WHERE user_id=?",(u["id"],)).fetchone(); c.close()
    return bool(p and p[name])

def role_is(*roles):
    u=user()
    return bool(u and u["role"] in roles)

@APP.route("/requests")
def requests_page():
    if not has_permission("view_requests"): return "غير مصرح",403
    q=request.args.get("q","").strip(); status=request.args.get("status",""); reference=request.args.get("reference",""); city=request.args.get("city","")
    c=db(); sql="""SELECT r.*,u.name creator,co.cr,co.unified,co.name company_name FROM requests r
                 LEFT JOIN users u ON u.id=r.created_by LEFT JOIN companies co ON co.id=r.company_id WHERE 1=1"""
    args=[]
    if q: sql+=" AND (r.serial LIKE ? OR r.request_no LIKE ? OR co.cr LIKE ? OR co.name LIKE ?)"; args += [f"%{q}%"]*4
    if status: sql+=" AND r.status=?"; args.append(status)
    if reference: sql+=" AND r.reference=?"; args.append(reference)
    if city: sql+=" AND r.city=?"; args.append(city)
    sql+=" ORDER BY r.id DESC"
    rows=c.execute(sql,args).fetchall(); c.close()
    return render_template("requests.html",title="الطلبات",rows=rows,q=q,status=status,reference=reference,city=city,references=REFERENCES,cities=CITIES)

@APP.route("/requests/new",methods=["GET","POST"])
def new_request():
    if not has_permission("create_request"): return "غير مصرح",403
    c=db(); companies=c.execute("SELECT * FROM companies WHERE active=1 ORDER BY name").fetchall()
    if request.method=="POST":
        serial="REQ-"+datetime.now().strftime("%Y%m%d")+"-"+secrets.token_hex(3).upper()
        c.execute("""INSERT INTO requests(serial,request_no,request_type,reference,city,company_id,created_by,created_at)
                     VALUES(?,?,?,?,?,?,?,?)""",(serial,request.form["request_no"],request.form["request_type"],request.form["reference"],request.form["city"],request.form["company_id"],session["uid"],now()))
        rid=c.execute("SELECT last_insert_rowid()").fetchone()[0]; c.commit(); c.close()
        audit("إنشاء طلب","request",rid,serial); return redirect(url_for("request_detail",rid=rid))
    c.close(); return render_template("new_request.html",title="طلب جديد",references=REFERENCES,cities=CITIES,companies=companies)

@APP.route("/requests/<int:rid>")
def request_detail(rid):
    c=db(); r=c.execute("""SELECT r.*,u.name creator,co.cr,co.unified,co.name company_name FROM requests r
                           LEFT JOIN users u ON u.id=r.created_by LEFT JOIN companies co ON co.id=r.company_id WHERE r.id=?""",(rid,)).fetchone()
    f=c.execute("SELECT f.*,u.name FROM followups f LEFT JOIN users u ON u.id=f.user_id WHERE f.request_id=? ORDER BY f.id DESC",(rid,)).fetchall()
    emails=c.execute("SELECT * FROM email_messages WHERE request_id=? ORDER BY received_at DESC, id DESC",(rid,)).fetchall()
    c.close()
    if not r: return "غير موجود",404
    return render_template("request_detail.html",title="تفاصيل الطلب",r=r,f=f,emails=emails)

@APP.route("/requests/<int:rid>/followup",methods=["POST"])
def followup(rid):
    c=db(); t=now(); note=request.form["note"].strip()
    c.execute("INSERT INTO followups(request_id,user_id,note,created_at) VALUES(?,?,?,?)",(rid,session["uid"],note,t))
    c.execute("UPDATE requests SET last_followup=?,status='قيد التنفيذ' WHERE id=?",(t,rid)); c.commit()
    r=c.execute("SELECT * FROM requests WHERE id=?",(rid,)).fetchone(); c.close()
    audit("إضافة متابعة","request",rid,note)
    notify_managers("followup",{"request_no":r["request_no"],"serial":r["serial"],"user":user()["name"],"date":t,"note":note})
    return redirect(url_for("request_detail",rid=rid))

@APP.route("/requests/<int:rid>/close-request",methods=["POST"])
def close_request(rid):
    c=db(); t=now(); note=request.form["note"].strip()
    c.execute("UPDATE requests SET status='بانتظار اعتماد الإنهاء',close_requested_at=?,close_note=? WHERE id=?",(t,note,rid)); c.commit()
    r=c.execute("SELECT * FROM requests WHERE id=?",(rid,)).fetchone(); c.close()
    audit("طلب اعتماد إنهاء","request",rid,note)
    notify_managers("close_request",{"request_no":r["request_no"],"user":user()["name"],"date":t,"note":note})
    return redirect(url_for("request_detail",rid=rid))

@APP.route("/requests/<int:rid>/approve",methods=["POST"])
def approve(rid):
    if user()["role"] not in ("manager","super_admin"): return "غير مصرح",403
    t=now(); c=db(); c.execute("UPDATE requests SET status='منتهي',closed_at=?,closed_by=? WHERE id=?",(t,session["uid"],rid)); c.commit()
    r=c.execute("SELECT r.*,u.email employee_email FROM requests r LEFT JOIN users u ON u.id=r.created_by WHERE r.id=?",(rid,)).fetchone(); c.close()
    audit("اعتماد إنهاء","request",rid,"تم اعتماد الإنهاء")
    notify_user(r["employee_email"],"closed",{"request_no":r["request_no"],"serial":r["serial"],"manager":user()["name"],"date":t})
    return redirect(url_for("request_detail",rid=rid))

@APP.route("/requests/<int:rid>/reject",methods=["POST"])
def reject(rid):
    if user()["role"] not in ("manager","super_admin"): return "غير مصرح",403
    note=request.form["note"].strip()
    c=db()
    c.execute("UPDATE requests SET status='قيد التنفيذ',close_requested_at=NULL,close_note=? WHERE id=?",(note,rid)); c.commit()
    r=c.execute("SELECT r.*,u.email employee_email FROM requests r LEFT JOIN users u ON u.id=r.created_by WHERE r.id=?",(rid,)).fetchone(); c.close()
    audit("رفض إنهاء","request",rid,note)
    notify_user(r["employee_email"],"rejected",{"request_no":r["request_no"],"serial":r["serial"],"manager":user()["name"],"date":now(),"note":note})
    return redirect(url_for("request_detail",rid=rid))

@APP.route("/requests/completed")
def completed_requests():
    if not has_permission("completed"): return "غير مصرح",403
    if not session.get("uid"): return redirect(url_for("login"))
    c=db()
    rows=c.execute("""SELECT r.*,u.name creator,co.cr,co.unified,co.name company_name
                      FROM requests r LEFT JOIN users u ON u.id=r.created_by
                      LEFT JOIN companies co ON co.id=r.company_id
                      WHERE r.status='منتهي' ORDER BY r.closed_at DESC, r.id DESC""").fetchall()
    c.close()
    return render_template("request_status.html",title="الطلبات المنتهية",heading="الطلبات المنتهية",sub="جميع الطلبات التي تم اعتماد إنهائها",rows=rows)

@APP.route("/requests/in-progress")
def in_progress_requests():
    if not has_permission("in_progress"): return "غير مصرح",403
    if not session.get("uid"): return redirect(url_for("login"))
    c=db()
    rows=c.execute("""SELECT r.*,u.name creator,co.cr,co.unified,co.name company_name
                      FROM requests r LEFT JOIN users u ON u.id=r.created_by
                      LEFT JOIN companies co ON co.id=r.company_id
                      WHERE r.status IN ('قيد التنفيذ','بانتظار اعتماد الإنهاء')
                      ORDER BY r.id DESC""").fetchall()
    c.close()
    return render_template("request_status.html",title="الطلبات قيد التنفيذ",heading="الطلبات قيد التنفيذ",sub="صفحة مستقلة للطلبات غير المنتهية",rows=rows)

@APP.route("/requests/10-days")
def ten_days_requests():
    if not has_permission("ten_days"): return "غير مصرح",403
    if not role_is("manager","super_admin"): return "غير مصرح",403
    cutoff=(datetime.now()-timedelta(days=10)).strftime("%Y-%m-%d %H:%M:%S")
    c=db()
    rows=c.execute("""SELECT r.*,u.name creator,co.cr,co.unified,co.name company_name
                      FROM requests r LEFT JOIN users u ON u.id=r.created_by
                      LEFT JOIN companies co ON co.id=r.company_id
                      WHERE r.status='قيد التنفيذ'
                        AND COALESCE(r.last_followup,r.created_at) <= ?
                      ORDER BY COALESCE(r.last_followup,r.created_at) ASC""",(cutoff,)).fetchall()
    c.close()
    return render_template("request_status.html",title="طلبات 10 أيام",heading="طلبات لم تتم متابعتها 10 أيام",sub="هذه الصفحة متاحة للمدراء وSuper Admin فقط",rows=rows)

@APP.route("/users",methods=["GET","POST"])
def users_page():
    if not has_permission("users_admin"): return "غير مصرح",403
    c=db()
    if request.method=="POST":
        c.execute("INSERT INTO users(name,username,password,email,role) VALUES(?,?,?,?,?)",(request.form["name"],request.form["username"],request.form["password"],request.form["email"],request.form["role"]))
        uid=c.execute("SELECT last_insert_rowid()").fetchone()[0]
        c.execute("INSERT INTO user_permissions(user_id,view_requests,create_request,completed,in_progress,audit,email_settings) VALUES(?,?,?,?,?,?,?)",(uid,1,1,1,1,1,0))
        c.commit(); audit("إضافة مستخدم","user",uid,request.form["username"])
    rows=c.execute("""SELECT u.*,p.view_requests,p.create_request,p.manage_companies,p.reports,p.completed,p.in_progress,p.ten_days,p.users_admin,p.email_templates,p.audit,p.email_integration,p.email_settings
                      FROM users u LEFT JOIN user_permissions p ON p.user_id=u.id ORDER BY u.id DESC""").fetchall(); c.close()
    return render_template("users.html",title="المستخدمون",rows=rows)

@APP.route("/users/<int:uid>/permissions",methods=["POST"])
def update_permissions(uid):
    if not has_permission("users_admin"): return "غير مصرح",403
    keys=["view_requests","create_request","manage_companies","reports","completed","in_progress","ten_days","users_admin","email_templates","audit","email_integration","email_settings"]
    vals=[1 if request.form.get(k)=="on" else 0 for k in keys]
    c=db(); c.execute("""INSERT INTO user_permissions(user_id,view_requests,create_request,manage_companies,reports,completed,in_progress,ten_days,users_admin,email_templates,audit,email_integration,email_settings)
      VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET
      view_requests=excluded.view_requests,create_request=excluded.create_request,manage_companies=excluded.manage_companies,reports=excluded.reports,completed=excluded.completed,in_progress=excluded.in_progress,ten_days=excluded.ten_days,users_admin=excluded.users_admin,email_templates=excluded.email_templates,audit=excluded.audit,email_integration=excluded.email_integration,email_settings=excluded.email_settings""",(uid,*vals)); c.commit(); c.close()
    audit("تعديل صلاحيات","user",uid,", ".join(k for k,v in zip(keys,vals) if v)); flash("تم حفظ الصلاحيات","success"); return redirect(url_for("users_page"))

@APP.route("/companies",methods=["GET","POST"])
def companies_page():
    if not has_permission("manage_companies"): return "غير مصرح",403
    c=db()
    if request.method=="POST":
        if request.form.get("action")=="upload":
            f=request.files.get("file")
            if f and f.filename:
                try:
                    imported=0
                    if f.filename.lower().endswith((".xlsx",".xlsm")):
                        wb=openpyxl.load_workbook(f,read_only=True,data_only=True); rows_data=list(wb.active.iter_rows(values_only=True))
                        headers=[str(x).strip() if x is not None else "" for x in (rows_data[0] if rows_data else [])]
                        records=[dict(zip(headers,row)) for row in rows_data[1:]]
                    else:
                        records=list(csv.DictReader(io.StringIO(f.read().decode("utf-8-sig"))))
                    for vals in records:
                        cr=vals.get("رقم السجل التجاري") or vals.get("cr") or vals.get("السجل")
                        unified=vals.get("الرقم الموحد") or vals.get("unified")
                        name=vals.get("اسم المنشأة") or vals.get("name")
                        if cr and unified and name:
                            cur=c.execute("INSERT OR IGNORE INTO companies(cr,unified,name) VALUES(?,?,?)",(str(cr).strip(),str(unified).strip(),str(name).strip())); imported+=cur.rowcount
                    c.commit(); audit("رفع سجلات تجارية","company",None,f"تم استيراد {imported} سجل"); flash(f"تم استيراد {imported} سجل بنجاح","success")
                except Exception as exc:
                    c.rollback(); flash(f"تعذر قراءة الملف: {exc}","error")
            else: flash("اختر ملف Excel أو CSV","error")
        else:
            c.execute("INSERT OR IGNORE INTO companies(cr,unified,name) VALUES(?,?,?)",(request.form["cr"],request.form["unified"],request.form["name"])); c.commit(); audit("إضافة سجل تجاري","company",None,request.form["cr"])
    rows=c.execute("SELECT * FROM companies ORDER BY id DESC").fetchall(); c.close()
    return render_template("companies.html",title="السجلات التجارية",rows=rows)

@APP.route("/companies/template")
def companies_template():
    if not has_permission("manage_companies"): return "غير مصرح",403
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="السجلات"
    ws.append(["رقم السجل التجاري","الرقم الموحد","اسم المنشأة"])
    ws.append(["1010XXXXXX","700XXXXXXX","اسم المنشأة"])
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,as_attachment=True,download_name="نموذج_رفع_السجلات.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@APP.route("/email-settings", methods=["GET","POST"])
def email_settings():
    if not has_permission("email_settings"): return "غير مصرح",403
    if request.method=="POST":
        env_path=os.path.join(os.path.dirname(__file__),".env")
        vals={k:request.form.get(k,"").strip() for k in ["MS_TENANT_ID","MS_CLIENT_ID","MS_CLIENT_SECRET","MS_SENDER_EMAIL","EMAIL_SYNC_SECONDS"]}
        lines=["# Microsoft 365 / Graph settings"]+[f"{k}={v}" for k,v in vals.items()]
        with open(env_path,"w",encoding="utf-8") as f: f.write("\n".join(lines)+"\n")
        os.environ.update(vals)
        audit("تحديث إعدادات ربط الإيميل","email_settings",None,"تم حفظ إعدادات Microsoft 365")
        flash("تم حفظ إعدادات ربط الإيميل. قد تحتاج لإعادة تشغيل النظام لتطبيقها بالكامل.","success")
    return render_template("email_settings.html",title="إعدادات ربط الإيميل",configured=is_configured(),values={k:os.environ.get(k,"") for k in ["MS_TENANT_ID","MS_CLIENT_ID","MS_CLIENT_SECRET","MS_SENDER_EMAIL","EMAIL_SYNC_SECONDS"]})
    
@APP.route("/audit")
def audit_page():
    if not has_permission("audit"): return "غير مصرح",403
    c=db(); rows=c.execute("""SELECT a.*,u.name FROM audit_logs a LEFT JOIN users u ON u.id=a.user_id ORDER BY a.id DESC LIMIT 300""").fetchall(); c.close()
    return render_template("audit.html",title="سجل التعديلات",rows=rows)

@APP.route("/templates",methods=["GET","POST"])
def templates_page():
    if not has_permission("email_templates"): return "غير مصرح",403
    if user()["role"]!="super_admin": return "غير مصرح",403
    c=db()
    if request.method=="POST":
        c.execute("UPDATE email_templates SET subject=?,body=? WHERE code=?",(request.form["subject"],request.form["body"],request.form["code"])); c.commit(); audit("تعديل قالب بريد","template",None,request.form["code"])
    rows=c.execute("SELECT * FROM email_templates").fetchall(); c.close()
    return render_template("templates.html",title="قوالب الإيميلات",rows=rows)

@APP.route("/reports", methods=["GET"])
def reports():
    if not has_permission("reports") or not role_is("hr","manager","super_admin"): return "غير مصرح",403
    return render_template("reports.html", title="التقارير")

@APP.route("/reports/export")
def reports_export():
    if not has_permission("reports") or not role_is("hr","manager","super_admin"): return "غير مصرح",403
    fmt=request.args.get("format","xlsx").lower()
    period=request.args.get("period","all")
    start=request.args.get("start","")
    end=request.args.get("end","")
    c=db()
    sql="""SELECT r.serial,r.request_no,r.request_type,r.reference,r.city,r.created_at,r.status,
                  r.last_followup,r.closed_at,co.cr,co.unified,co.name company_name
           FROM requests r LEFT JOIN companies co ON co.id=r.company_id WHERE 1=1"""
    args=[]
    if period=="custom":
        if start: sql+=" AND date(r.created_at)>=date(?)"; args.append(start)
        if end: sql+=" AND date(r.created_at)<=date(?)"; args.append(end)
    sql+=" ORDER BY r.id DESC"
    rows=c.execute(sql,args).fetchall(); c.close()
    headers=["الرقم التسلسلي","رقم الطلب","نوع الطلب","المرجع","المدينة","تاريخ الطلب","الحالة","آخر متابعة","تاريخ الإنهاء","السجل التجاري","الرقم الموحد","اسم المنشأة"]
    data=[[r["serial"],r["request_no"],r["request_type"],r["reference"],r["city"],r["created_at"],r["status"],r["last_followup"] or "",r["closed_at"] or "",r["cr"] or "",r["unified"] or "",r["company_name"] or ""] for r in rows]
    if fmt=="xlsx":
        wb=openpyxl.Workbook(); ws=wb.active; ws.title="تقرير الطلبات"
        ws.append(headers)
        for row in data: ws.append(row)
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=openpyxl.styles.Font(bold=True)
        for col in ws.columns:
            letter=col[0].column_letter
            ws.column_dimensions[letter].width=min(max(max(len(str(x.value or "")) for x in col)+2,12),35)
        out=io.BytesIO(); wb.save(out); out.seek(0)
        return send_file(out,as_attachment=True,download_name="تقرير_الطلبات.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if fmt=="pdf":
        out=io.BytesIO()
        doc=SimpleDocTemplate(out,pagesize=landscape(A4),rightMargin=18,leftMargin=18,topMargin=18,bottomMargin=18)
        styles=getSampleStyleSheet(); title=Paragraph("تقرير الطلبات - BFM",styles["Title"])
        table=Table([headers]+data,repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1f2937")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.4,colors.grey),("FONTSIZE",(0,0),(-1,-1),6),("VALIGN",(0,0),(-1,-1),"MIDDLE")]))
        doc.build([title,Spacer(1,8),table]); out.seek(0)
        return send_file(out,as_attachment=True,download_name="تقرير_الطلبات.pdf",mimetype="application/pdf")
    return "صيغة غير مدعومة",400

def sync_inbox():
    messages,error=fetch_inbox(50)
    if error: return False,error
    c=db(); added=0
    for m in messages:
        mid=m.get("internetMessageId") or m.get("id")
        if not mid or c.execute("SELECT 1 FROM email_messages WHERE message_id=?",(mid,)).fetchone(): continue
        subject=m.get("subject") or ""
        body=((m.get("body") or {}).get("content") or "")
        match=re.search(r"REQ:([A-Z0-9_-]+)",subject+"\n"+body,re.I)
        if not match: continue
        req=c.execute("SELECT id FROM requests WHERE serial=?",(match.group(1),)).fetchone()
        if not req: continue
        sender=((m.get("from") or {}).get("emailAddress") or {}).get("address","")
        recipients=", ".join(((x.get("emailAddress") or {}).get("address","")) for x in (m.get("toRecipients") or []))
        c.execute("""INSERT OR IGNORE INTO email_messages
          (request_id,direction,sender,recipients,subject,body,message_id,received_at)
          VALUES(?,?,?,?,?,?,?,?)""",(req["id"],"inbound",sender,recipients,subject,body,mid,m.get("receivedDateTime") or now()))
        added+=c.execute("SELECT changes()").fetchone()[0]
    c.commit(); c.close()
    return True,f"تمت مزامنة {added} رسالة"

def _email_sync_loop():
    while True:
        try:
            if is_configured(): sync_inbox()
        except Exception: pass
        time.sleep(max(int(os.environ.get("EMAIL_SYNC_SECONDS","60")),30))

init_db()
if is_configured():
    threading.Thread(target=_email_sync_loop,daemon=True,name="email-sync").start()
if __name__=="__main__":
    APP.run(host="0.0.0.0",port=5000,debug=False)
